#!/usr/bin/env python3
"""
campo_chars.py -- Cuenta caracteres por campo del pautado.

Dos diferencias con un conteo global:
1. Cada campo tiene su propio limite; el total puede pasar aunque un campo
   este bien, o al reves.
2. Los limites se leen de la plantilla oficial cuando la plantilla los declara
   ("Problema (max. 1,250 caracteres)"). Si no, se usan los limites vigentes.

Cuenta con espacios, igual que la maqueta del certamen.
"""
from __future__ import annotations

import pathlib
import re
import sys

# Limites vigentes con espacios. Fuente: encabezados de la maqueta oficial.
LIMITS: dict[str, int] = {
    "problema": 1250,
    "solucion": 1250,
    "prefactibilidad": 1000,
    "innovacion": 1000,
    "impacto": 1500,
}

_ALIAS = {
    "problema": "problema",
    "solución": "solucion",
    "solucion": "solucion",
    "prefactibilidad": "prefactibilidad",
    "innovación": "innovacion",
    "innovacion": "innovacion",
    "impacto": "impacto",
}
_CAMPO_RE = re.compile(
    r"(problema|soluci[oó]n|prefactibilidad|innovaci[oó]n|impacto)"
    r"[^\n]{0,90}?"
    r"(\d[\d.,]{2,})\s*(?:caracteres|chars)",
    re.IGNORECASE,
)
_MARCADOR_RE = re.compile(r"\{SEC_([A-Z_]+)\}(.*?)\{SEC_END\}", re.DOTALL | re.IGNORECASE)


def _normalizar(nombre: str) -> str:
    return _ALIAS.get(nombre.strip().lower(), nombre.strip().lower())


def _entero(token: str) -> int:
    return int(re.sub(r"[.,]", "", token))


def limites_desde_plantilla(ruta: pathlib.Path) -> dict[str, int]:
    """Lee los limites declarados en la plantilla. Devuelve {} si no declara."""
    if not ruta.exists():
        return {}
    try:
        import docx

        doc = docx.Document(str(ruta))
    except Exception:
        return {}
    partes = [p.text for p in doc.paragraphs]
    for tabla in doc.tables:
        for fila in tabla.rows:
            partes.extend(celda.text for celda in fila.cells)
    texto = "\n".join(partes)
    limites: dict[str, int] = {}
    for match in _CAMPO_RE.finditer(texto):
        campo = _normalizar(match.group(1))
        limites[campo] = _entero(match.group(2))
    return limites


def texto_docx(ruta: pathlib.Path) -> str:
    import docx

    doc = docx.Document(str(ruta))
    partes = [p.text for p in doc.paragraphs]
    for tabla in doc.tables:
        for fila in tabla.rows:
            partes.extend(celda.text for celda in fila.cells)
    return "\n".join(partes)


def secciones_por_marcador(texto: str) -> dict[str, str]:
    return {
        _normalizar(match.group(1)): match.group(2).strip()
        for match in _MARCADOR_RE.finditer(texto)
    }


def secciones_por_encabezado(texto: str) -> dict[str, str]:
    """Fallback: acumula parrafos bajo el encabezado de cada campo."""
    secciones: dict[str, list[str]] = {}
    actual: str | None = None
    for linea in texto.splitlines():
        limpia = linea.strip()
        if not limpia:
            continue
        candidato = _normalizar(re.sub(r"[:.].*$", "", limpia))
        if candidato in LIMITS and len(limpia) <= 60:
            actual = candidato
            secciones.setdefault(actual, [])
            continue
        if actual:
            secciones[actual].append(limpia)
    return {campo: " ".join(lineas) for campo, lineas in secciones.items()}


def contar(texto: str) -> dict[str, int]:
    secciones = secciones_por_marcador(texto)
    if not secciones:
        secciones = secciones_por_encabezado(texto)
    return {campo: len(valor.strip()) for campo, valor in secciones.items()}


def revisar(docx_path: pathlib.Path, plantilla: pathlib.Path | None = None) -> tuple[list[dict], dict[str, int]]:
    plantilla = plantilla or pathlib.Path("docs/pautado_template.docx")
    limites = {**LIMITS, **limites_desde_plantilla(plantilla)}
    texto = texto_docx(docx_path)
    conteos = contar(texto)
    hallazgos: list[dict] = []
    if not conteos:
        hallazgos.append({
            "gravedad": "critico",
            "que": "No se reconocio ningun campo del pautado",
            "hacer": "Conserva los marcadores {SEC_CAMPO}...{SEC_END} o los encabezados oficiales",
        })
        return hallazgos, conteos
    for campo, limite in limites.items():
        if campo not in conteos:
            hallazgos.append({
                "gravedad": "importante",
                "campo": campo,
                "que": f"Campo ausente (limite {limite})",
                "hacer": "Agrega el campo o corrige el marcador antes de exportar",
            })
            continue
        cuenta = conteos[campo]
        if cuenta > limite:
            hallazgos.append({
                "gravedad": "critico",
                "campo": campo,
                "que": f"{cuenta} / {limite} caracteres ({cuenta - limite} de exceso)",
                "hacer": f"Recorta {cuenta - limite} caracteres de {campo}",
            })
    return hallazgos, conteos


def main(argv: list[str] | None = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    docx_path = pathlib.Path(argv[0]) if argv else pathlib.Path("docs/pautado_draft.docx")
    plantilla = pathlib.Path(argv[1]) if len(argv) > 1 else pathlib.Path("docs/pautado_template.docx")
    if not docx_path.exists():
        print(f"  PENDIENTE: {docx_path} no existe. Ejecuta python scripts/gen_pautado.py primero.")
        return 1
    hallazgos, conteos = revisar(docx_path, plantilla)
    print("  CAMPO            CUENTA / LIMITE")
    for campo, limite in LIMITS.items():
        cuenta = conteos.get(campo)
        if cuenta is None:
            print(f"  {campo.upper():<16} {'--':>6} / {limite:<6}  ausente")
            continue
        marca = "X" if cuenta > limite else "v"
        print(f"  [{marca}] {campo.upper():<13} {cuenta:>6} / {limite:<6}  {(cuenta / limite) * 100:5.1f}%")
    if hallazgos:
        print()
        for h in hallazgos:
            icono = "X" if h["gravedad"] == "critico" else "!"
            print(f"  [{icono}] {h.get('campo', '--')}: {h['que']} -> {h['hacer']}")
        return 1 if any(h["gravedad"] == "critico" for h in hallazgos) else 0
    print("\n  OK: todos los campos respetan su limite.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
