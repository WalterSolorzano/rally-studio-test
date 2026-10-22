# -*- coding: utf-8 -*-
"""Sincroniza el corpus del cerebro (brain/corpus) desde las fuentes canonicas.

Regla: el cerebro aprende SOLO de conocimiento de dominio (rubricas, ganadores,
retos, claims, metodologia). Nunca del motor, ni de proyectos anteriores.

Uso:
    python brain/sync_corpus.py

Despues de sincronizar, reconstruir el grafo:
    graphify-out/.graphify_python  (ver brain/README.md)
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CORPUS = ROOT / "brain" / "corpus"

# Fuente -> nombre limpio dentro del corpus.
# Nombres en ASCII a proposito: el graphify deriva los IDs de nodo de la ruta.
COPY = {
    "knowledge/winners_latam.md": "ganadores_latam.md",
    "knowledge/winners_analysis.md": "ganadores_analisis.md",
    "knowledge/rally_analysis.txt": "rally_analisis.txt",
    "knowledge/rubric_matrices.md": "rubricas_matrices.md",
    "knowledge/rubric_latam.md": "rubrica_latam.md",
    "knowledge/voice_notes.md": "guion_voz.md",
    "research/challenges_report.md": "retos_reporte.md",
    # NOTA: "Proyectos Ganadores de Rally Latam.md" y
    # "Análisis del Rally de Innovación.txt" de la raiz son duplicados byte a byte
    # de knowledge/winners_latam.md y knowledge/rally_analysis.txt.
    # Se usa la version de knowledge/ (la canonica segun knowledge/INDICE.md).
    "chatgptdeep-research-report.md": "investigacion_pitch.md",
    "Rally_Innovacion_Nicaragua_Analisis.pdf": "analisis_nicaragua.pdf",
    "rubrica_oficial_latam.jpeg": "rubrica_oficial.jpeg",
}

# JSON -> Markdown legible (si van como .json, graphify los trata como codigo/AST).
JSON_DOCS = {
    "research/claims.json": ("claims_datos.md", "Claims verificables"),
    "knowledge/winners_database.json": ("ganadores_db.md", "Casos ganadores por fase"),
    "research/challenges.json": ("retos.md", "Banco de retos"),
    "research/challenges_ranked.json": ("retos_ranked.md", "Retos priorizados"),
}

# DOCX -> Markdown (texto plano, sin formato). Mas determinista para extraer.
DOCX = {
    "knowledge/pautado_nacional_ref.docx": "pautado_nacional_ref.md",
    "knowledge/pautado_latam_ref.docx": "pautado_latam_ref.md",
}

# Nunca entran al cerebro.
# - LEEME.md / docs/ / scripts/ / dashboard/: instrucciones y motor, no conocimiento.
# - graphify-out/converted/: proyectos anteriores (regla 11 de AGENTS.md).
# - imagenes sueltas sin valor probatorio.


def docx_to_md(src: Path) -> str:
    from docx import Document

    doc = Document(str(src))
    out: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(c for c in style if c.isdigit()) or "2"
            out.append("#" * min(int(level) + 1, 6) + " " + t)
        else:
            out.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            out.append("| " + " | ".join(cells) + " |")
    return "\n\n".join(out)


def json_to_md(src: Path, title: str) -> str:
    """Renderiza datos JSON como Markdown legible.

    Sin esto, graphify clasifica los .json como codigo y los pasa por AST,
    que produce nodos basura. Como documento, el extractor semantico si
    entiende los claims y los casos ganadores.
    """
    data = json.loads(src.read_text(encoding="utf-8"))
    out: list[str] = [f"# {title}", ""]

    def scalar(v):
        if isinstance(v, bool):
            return "si" if v else "no"
        if v is None:
            return ""
        if isinstance(v, (int, float)):
            return str(v)
        return str(v).replace("\n", " ").strip()

    def walk(obj, depth=0):
        pad = "  " * depth
        if isinstance(obj, dict):
            for k, v in obj.items():
                key = str(k).replace("_", " ").strip()
                if isinstance(v, (dict, list)):
                    out.append("")
                    out.append(f"{'#' * min(depth + 2, 6)} {key}")
                    out.append("")
                    walk(v, depth + 1)
                else:
                    out.append(f"- **{key}**: {scalar(v)}")
        elif isinstance(obj, list):
            for i, item in enumerate(obj, 1):
                if isinstance(item, (dict, list)):
                    out.append("")
                    out.append(f"{'#' * min(depth + 2, 6)} Registro {i}")
                    out.append("")
                    walk(item, depth + 1)
                else:
                    out.append(f"- {scalar(item)}")
        else:
            out.append(f"- {scalar(obj)}")

    walk(data)
    return "\n".join(out)


def main() -> int:
    CORPUS.mkdir(parents=True, exist_ok=True)
    written, missing = [], []

    for rel, name in COPY.items():
        src = ROOT / rel
        if not src.is_file():
            missing.append(rel)
            continue
        dst = CORPUS / name
        shutil.copyfile(src, dst)
        written.append((name, dst.stat().st_size))

    for rel, name in DOCX.items():
        src = ROOT / rel
        if not src.is_file():
            missing.append(rel)
            continue
        text = docx_to_md(src)
        dst = CORPUS / name
        dst.write_text(text, encoding="utf-8")
        written.append((name, dst.stat().st_size))

    for rel, (name, title) in JSON_DOCS.items():
        src = ROOT / rel
        if not src.is_file():
            missing.append(rel)
            continue
        text = json_to_md(src, title)
        dst = CORPUS / name
        dst.write_text(text, encoding="utf-8")
        written.append((name, dst.stat().st_size))

    # El proyecto real entra como un archivo mas, y solo cuando exista.
    # sync_corpus.py no lo borra: si esta, se queda.
    project = CORPUS / "proyecto_actual.md"
    state = {
        "corpus_files": len(list(CORPUS.glob("*"))),
        "project_file": project.name if project.is_file() else None,
        "missing_sources": missing,
        "files": [{"name": n, "bytes": s} for n, s in sorted(written)],
    }
    (ROOT / "brain" / "corpus_state.json").write_text(
        json.dumps(state, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    print(f"Corpus: {len(written)} archivos -> {CORPUS}")
    print(f"Proyecto actual: {project.name if project.is_file() else 'NO (el cerebro espera)'}")
    if missing:
        print("Fuentes ausentes: " + ", ".join(missing))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
