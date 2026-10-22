#!/usr/bin/env python3
"""
smoke.py -- Prueba de humo del arsenal.

Verifica en segundos que las piezas que pueden costar el certamen funcionan en
esta maquina: lectura de DOCX, YAML, medicion real de duracion con ffmpeg,
conteo por campo y cruce de cifras. No necesita GPU ni red.
"""
from __future__ import annotations

import pathlib
import struct
import sys
import tempfile
import wave

RAIZ = pathlib.Path(__file__).resolve().parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from checks import campo_chars, coherencia, duracion  # noqa: E402

RESULTADOS: list[tuple[str, bool, str]] = []


def probar(nombre: str, funcion) -> None:
    try:
        detalle = funcion()
        RESULTADOS.append((nombre, True, str(detalle)))
    except Exception as exc:  # noqa: BLE001
        RESULTADOS.append((nombre, False, f"{type(exc).__name__}: {exc}"))


def _wav_de_un_segundo(ruta: pathlib.Path) -> pathlib.Path:
    with wave.open(str(ruta), "w") as f:
        f.setnchannels(1)
        f.setsampwidth(2)
        f.setframerate(16000)
        for i in range(16000):
            f.writeframes(struct.pack("<h", int(3000 * (i % 100 < 50))))


def prueba_duracion() -> str:
    with tempfile.TemporaryDirectory() as tmp:
        ruta = pathlib.Path(tmp) / "tono.wav"
        _wav_de_un_segundo(ruta)
        segundos, _, motor = duracion.medir(ruta)
        assert segundos is not None, "no se pudo medir un WAV de 1 s"
        assert 0.8 <= segundos <= 1.3, f"duracion inesperada: {segundos}"
        return f"{segundos:.2f} s via {motor}"


def prueba_campo() -> str:
    import docx

    with tempfile.TemporaryDirectory() as tmp:
        ruta = pathlib.Path(tmp) / "pautado.docx"
        doc = docx.Document()
        doc.add_paragraph("{SEC_PROBLEMA}Texto corto{/SEC_END}".replace("{/SEC_END}", "{SEC_END}"))
        doc.add_paragraph("{SEC_IMPACTO}" + ("x" * 1600) + "{SEC_END}")
        doc.save(str(ruta))
        hallazgos, conteos = campo_chars.revisar(ruta, pathlib.Path("no_existe.docx"))
        assert conteos["problema"] == len("Texto corto"), conteos
        assert any(h["campo"] == "impacto" for h in hallazgos), hallazgos
        return "marcadores y exceso detectados"


def prueba_coherencia() -> str:
    a = coherencia.desde_texto("Ahorra 1,250 litros y cubre el 64% de la demanda.")
    b = coherencia.desde_texto("Ahorra 1250 litros y cubre el 46% de la demanda.")
    hallazgos = coherencia.comparar({"guion": a, "pautado": b})
    numeros = {h["numero"] for h in hallazgos}
    assert "64" in numeros and "46" in numeros, numeros
    assert "1250" not in numeros, "1,250 y 1250 deben normalizarse igual"
    return f"{len(hallazgos)} discrepancias, miles normalizados"


def prueba_lock() -> str:
    import yaml

    datos = yaml.safe_load((RAIZ / "arsenal.lock.yaml").read_text(encoding="utf-8"))
    assert datos["version"] == 1
    assert any(e.get("nombre") == "ffmpeg" for e in datos["video"])
    return "arsenal.lock.yaml legible"


def prueba_brief() -> str:
    import json

    brief = json.loads((RAIZ / "brief.json").read_text(encoding="utf-8"))
    assert "rally_type" in brief and "project_name" in brief
    return "brief.json legible (vacio es valido: el cerebro espera)"


def main() -> int:
    print("SMOKE -- arsenal de Rally Studio\n")
    probar("ffmpeg localizado", lambda: duracion.find_ffmpeg() or (_ for _ in ()).throw(AssertionError("no hay ffmpeg")))
    probar("medicion de duracion", prueba_duracion)
    probar("caracteres por campo", prueba_campo)
    probar("coherencia de cifras", prueba_coherencia)
    probar("registro arsenal", prueba_lock)
    probar("brief", prueba_brief)

    for nombre, ok, detalle in RESULTADOS:
        print(f"  [{'PASA' if ok else 'FALLA'}] {nombre:<24} {detalle}")
    fallas = [r for r in RESULTADOS if not r[1]]
    print(f"\n  {len(RESULTADOS) - len(fallas)}/{len(RESULTADOS)} pruebas pasan.")
    return 1 if fallas else 0


if __name__ == "__main__":
    raise SystemExit(main())
